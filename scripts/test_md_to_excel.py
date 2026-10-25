"""
test_md_to_excel.py
────────────────────────────────────────────────────────────────────────────
Script de prueba: markdown OCR → JSON → Excel usando DeepSeek.

Reutiliza los archivos *__mistral.md de pruebas/ (ya generados por Mistral)
SIN hacer nuevas llamadas OCR. Solo llama a DeepSeek para estructurar.

Uso:
    cd "h:\git\OCR con Glm\OCR-con-GLM-IA"
    .venv\Scripts\python.exe scripts\test_md_to_excel.py

Resultado: pruebas\output\facturas_<timestamp>.json  y  .xlsx
"""
from __future__ import annotations

import json
import re
import sys
import logging
from datetime import datetime
from pathlib import Path

# ── Añadir raíz del proyecto al path ────────────────────────────────────────
ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    datefmt="%H:%M:%S",
)
_log = logging.getLogger("md2excel")

# ── CONFIGURACIÓN ────────────────────────────────────────────────────────────
# Seguridad: la API key NUNCA va en el código (una clave filtrada fue eliminada
# del historial con git-filter-repo). Léela del entorno o del .env del proyecto;
# en producción la capa SaaS la obtiene vía app.core.config.get_secret().
import os

DEEPSEEK_API_KEY = os.environ.get("DEEPSEEK_API_KEY", "")
PROVIDER        = "DeepSeek"
MODEL           = "deepseek-v4-flash"
MD_FOLDER       = ROOT / "pruebas"
OUTPUT_FOLDER   = MD_FOLDER / "output"
TIMESTAMP_RE    = re.compile(r"_\d{8}_\d{6}__mistral$")

# Columnas exactas en el orden esperado por el usuario
COLUMNS = [
    "Nombre archivo",
    "Número Factura",
    "Fecha",
    "Nombre proveedor",
    "N.I.F.",
    "Conceptos",
    "Base",
    "IVA",
    "IRPF %",
    "Importe",
    "B Exenta",
    "Total",
    "Alertas",
]

# ── PROMPT del usuario ────────────────────────────────────────────────────────
USER_PROMPT = """Eres un experto en extracción de facturas (España / Portugal). Recibes como entrada el texto extraído por OCR entre los marcadores ===BEGIN_OCR=== y ===END_OCR===. Analiza el texto y devuelve ÚNICAMENTE un JSON válido (sin texto adicional) con un array de objetos, uno por cada factura encontrada. Cada objeto debe tener exactamente estas claves (nombres en español, respetar mayúsculas tal como abajo):

[
{
"Nombre archivo": "", // string — nombre del fichero original (si no está, devolver "")
"Número Factura": "", // string — tal cual aparece
"Fecha": "", // string — ISO YYYY-MM-DD si es posible, si no transcribe tal cual
"Nombre proveedor": "", // string
"N.I.F.": "", // string
"Conceptos": "", // string — descripciones concatenadas separadas por " | "
"Base": 0.0, // float — suma de bases imponibles
"IVA": 0.0, // float — suma de importes IVA
"IRPF %": 0.0, // float — porcentaje de IRPF (numero sin %). 0.0 si no aplica
"Importe": 0.0, // float — base + IVA (antes de IRPF)
"B Exenta": 0.0, // float — base exenta
"Total": 0.0 // float — total final (aplicar IRPF si corresponde)
}
]

Reglas y prioridad de extracción:

Extrae Número Factura completo tal como aparece (no formatear).
Fecha: preferir ISO YYYY-MM-DD; si aparece DD/MM/YY o DD/MM/YYYY convertir; si solo mes/año, poner YYYY-MM-01; si no se puede, dejar el texto original.
Nombre proveedor y N.I.F.: identificar al emisor (no al receptor). El emisor suele aparecer en el encabezado o membrete de la factura; el receptor aparece en secciones como "DATOS DEL CLIENTE", "FACTURAR A", "CLIENTE", "A/ATT:" o similar. Si hay duda, elegir el nombre que aparezca en la parte superior del documento o junto al CIF/NIF del vendedor.
Conceptos: concatenar las líneas de detalle relevantes (líneas de descripción), separadas por " | ". Truncar a 300 chars si es excesivo.
Base, IVA, B Exenta, Total: devolver en euros como float con punto decimal. Si en OCR usan coma decimal, convertir a punto. Si moneda distinta, incluir nota de moneda en Conceptos y convertir los números tal cual aparecen.
IRPF %: si aparece retención (p.ej. "Retención 15%"), devolver 15.0; si no aparece, 0.0.
Importe: calcular como Base + IVA si no aparece explícitamente.
Total: debe reflejar el importe final pagadero (si hay retención: Total = Base + IVA - (Base * IRPF_pct/100)). Si el documento muestra un Total diferente por descuentos o ajustes, use el Total explícito del documento.
null solo si el campo es completamente inexistente e ilegible; para campos numéricos preferir 0.0 si es plausible inferir que no existe.
Reglas de formato: todos los campos numéricos deben ser números (no strings). Las claves deben aparecer exactamente como en el esquema.

Alertas: campo string obligatorio. Dejar "" si todo es correcto. Si detectas alguno de estos problemas, describir brevemente separados por " | ":
- Texto OCR ilegible o muy degradado en alguna zona relevante
- No se puede distinguir claramente el emisor del receptor
- Importe total no cuadra con Base + IVA - retención
- Número de factura ausente o ambiguo
- Fecha ausente o no convertible a formato estándar
- NIF/CIF con formato no reconocido o ausente
- Datos contradictorios entre sí en el documento
- Moneda distinta de euro
- Documento que no parece una factura (recibo, proforma, ticket sin datos fiscales)

Salida exigida:

Sólo el JSON array, sin texto adicional, sin bloques markdown. Ejemplo mínimo:
[
{
"Nombre archivo":"mi_factura.pdf",
"Número Factura":"FT-1234",
"Fecha":"2026-04-01",
"Nombre proveedor":"ACME S.L.",
"N.I.F.":"B12345678",
"Conceptos":"Servicio X | Producto Y",
"Base":100.0,
"IVA":21.0,
"IRPF %":0.0,
"Importe":121.0,
"B Exenta":0.0,
"Total":121.0,
"Alertas":""
}
]
Notas para el motor:

Si el OCR incluye varias facturas en el mismo PDF, devolver un objeto por cada factura.
No añadas campos distintos a los especificados.
No incluyas explicaciones, tablas ni formatos alternativos."""


# ── Helpers ───────────────────────────────────────────────────────────────────

def _filename_from_stem(stem: str) -> str:
    """Extrae el nombre original eliminando el sufijo de timestamp __mistral."""
    return TIMESTAMP_RE.sub("", stem).strip()


def _parse_json_response(raw: str) -> list[dict]:
    """Extrae el array JSON de la respuesta del LLM."""
    raw = raw.strip()
    # Eliminar bloques markdown si los hay
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.IGNORECASE)
    raw = re.sub(r"\s*```$", "", raw)
    raw = raw.strip()
    try:
        parsed = json.loads(raw)
        if isinstance(parsed, list):
            return parsed
        if isinstance(parsed, dict):
            return [parsed]
    except json.JSONDecodeError:
        # Intentar extraer el primer array JSON del texto
        m = re.search(r"\[[\s\S]*\]", raw)
        if m:
            try:
                parsed = json.loads(m.group(0))
                if isinstance(parsed, list):
                    return parsed
            except Exception:
                pass
        # Intentar extraer objeto suelto
        m2 = re.search(r"\{[\s\S]*\}", raw)
        if m2:
            try:
                parsed = json.loads(m2.group(0))
                if isinstance(parsed, dict):
                    return [parsed]
            except Exception:
                pass
    return []


def _normalize_row(row: dict, nombre_archivo: str) -> dict:
    """Asegura que el row tenga exactamente las columnas esperadas."""
    out: dict = {}
    for col in COLUMNS:
        val = row.get(col)
        if col == "Nombre archivo":
            out[col] = nombre_archivo if not val else str(val)
        elif col in ("Base", "IVA", "IRPF %", "Importe", "B Exenta", "Total"):
            try:
                out[col] = float(val) if val is not None else 0.0
            except (TypeError, ValueError):
                out[col] = 0.0
        elif col == "Alertas":
            out[col] = str(val).strip() if val else ""
        else:
            out[col] = str(val) if val is not None else ""
    return out


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    # Importar el motor de estructuración
    try:
        from src.core.structuring_engine import StructuringEngine
    except ImportError as exc:
        _log.error("No se pudo importar StructuringEngine: %s", exc)
        _log.error("Asegúrate de ejecutar desde la raíz del proyecto con el venv activado.")
        sys.exit(1)

    try:
        import pandas as pd
    except ImportError:
        _log.error("pandas no instalado. Ejecuta: pip install pandas openpyxl")
        sys.exit(1)

    # Inicializar engine
    try:
        engine = StructuringEngine(
            provider=PROVIDER,
            api_key=DEEPSEEK_API_KEY,
            model=MODEL,
        )
        _log.info("Engine inicializado: proveedor=%s modelo=%s", PROVIDER, MODEL)
    except Exception as exc:
        _log.error("Error inicializando StructuringEngine: %s", exc)
        sys.exit(1)

    # Buscar archivos markdown
    md_files = sorted(MD_FOLDER.glob("*__mistral.md"))
    if not md_files:
        _log.error("No se encontraron archivos *__mistral.md en %s", MD_FOLDER)
        sys.exit(1)

    _log.info("Encontrados %d archivos markdown en %s", len(md_files), MD_FOLDER)

    all_rows: list[dict] = []
    errores: list[str] = []

    for md_path in md_files:
        nombre_original = _filename_from_stem(md_path.stem)
        _log.info("── Procesando: %s", md_path.name)

        # Leer contenido markdown
        try:
            ocr_text = md_path.read_text(encoding="utf-8")
        except Exception as exc:
            _log.error("  No se pudo leer %s: %s", md_path.name, exc)
            errores.append(md_path.name)
            continue

        if not ocr_text.strip():
            _log.warning("  Archivo vacío: %s", md_path.name)
            errores.append(md_path.name)
            continue

        # Llamar a DeepSeek
        try:
            raw_json, pt, ct = engine.structure(ocr_text, USER_PROMPT)
            _log.info("  → Respuesta recibida (prompt_tokens=%s, completion_tokens=%s)", pt, ct)
        except Exception as exc:
            _log.error("  Error en structuring: %s", exc)
            errores.append(md_path.name)
            continue

        if not raw_json or not raw_json.strip():
            _log.warning("  LLM devolvió respuesta vacía para %s", md_path.name)
            errores.append(md_path.name)
            continue

        # Parsear JSON
        rows = _parse_json_response(raw_json)
        if not rows:
            _log.warning("  No se pudo parsear JSON para %s. Raw: %s", md_path.name, raw_json[:300])
            errores.append(md_path.name)
            continue

        _log.info("  → %d factura(s) extraída(s)", len(rows))

        for row in rows:
            normalized = _normalize_row(row, nombre_original)
            all_rows.append(normalized)

    _log.info("─" * 60)
    _log.info("Total filas: %d  |  Archivos con error: %d", len(all_rows), len(errores))

    if not all_rows:
        _log.error("No se extrajo ninguna fila. Revisa los logs anteriores.")
        if errores:
            _log.error("Archivos con error: %s", ", ".join(errores))
        sys.exit(1)

    # Crear carpeta de salida
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")

    # Guardar JSON
    json_path = OUTPUT_FOLDER / f"facturas_{ts}.json"
    json_path.write_text(
        json.dumps(all_rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    _log.info("JSON guardado: %s", json_path)

    # Guardar Excel
    xlsx_path = OUTPUT_FOLDER / f"facturas_{ts}.xlsx"
    df = pd.DataFrame(all_rows, columns=COLUMNS)

    with pd.ExcelWriter(str(xlsx_path), engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Facturas")
        ws = writer.sheets["Facturas"]

        # Ajuste básico de ancho de columnas
        col_widths = {
            "Nombre archivo": 22,
            "Número Factura": 18,
            "Fecha": 14,
            "Nombre proveedor": 40,
            "N.I.F.": 14,
            "Conceptos": 55,
            "Base": 12,
            "IVA": 10,
            "IRPF %": 10,
            "Importe": 12,
            "B Exenta": 12,
            "Total": 12,
            "Alertas": 45,
        }
        for i, col in enumerate(COLUMNS, 1):
            ws.column_dimensions[
                ws.cell(row=1, column=i).column_letter
            ].width = col_widths.get(col, 14)

        # Colorear en rojo las filas con Alertas no vacío
        try:
            from openpyxl.styles import PatternFill, Font
            red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
            red_font = Font(color="CC0000", bold=True)
            alertas_col_idx = COLUMNS.index("Alertas") + 1  # 1-based
            for row_idx, row in enumerate(all_rows, start=2):  # fila 1 = cabecera
                if row.get("Alertas", "").strip():
                    for col_idx in range(1, len(COLUMNS) + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        cell.fill = red_fill
                    # La celda de Alertas también en rojo oscuro
                    ws.cell(row=row_idx, column=alertas_col_idx).font = red_font
        except Exception as _exc:
            _log.warning("No se pudo aplicar color rojo en Excel: %s", _exc)

    _log.info("Excel guardado: %s", xlsx_path)
    _log.info("=" * 60)
    _log.info("LISTO. Abre el Excel en: %s", xlsx_path)

    if errores:
        _log.warning("Archivos que fallaron: %s", ", ".join(errores))


if __name__ == "__main__":
    main()
