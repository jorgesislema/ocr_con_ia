import sys
from pathlib import Path

sys.path.insert(0, ".")

from src.core.excel_exporter import export_invoices_to_excel


def _sample_invoices():
    return [
        {
            "invoice_id": "INV-0001",
            "fecha": "2026-04-01",
            "proveedor": "ACME S.A.",
            "moneda": "USD",
            "subtotal": "100.00",
            "iva": "12.00",
            "total": "112.00",
            "lines": [
                {
                    "origen": "Factura 0001",
                    "descripcion": "Producto A",
                    "cantidad": "1",
                    "precio_unitario": "$100.00",
                    "iva_pct": "12%",
                    "subtotal": "100.00",
                    "iva_amount": "12.00",
                    "total": "112.00",
                }
            ],
        },
        {
            "invoice_id": "INV-0002",
            "fecha": "01/04/2026",
            "proveedor": "BETA LTD",
            "moneda": "USD",
            "subtotal": "200,00",
            "iva": "24,00",
            "total": "224,00",
            "lines": [
                {
                    "origen": "Factura 0002",
                    "descripcion": "Servicio B",
                    "cantidad": "2",
                    "precio_unitario": "100,00",
                    "iva_pct": "12%",
                    "subtotal": "200,00",
                    "iva_amount": "24,00",
                    "total": "224,00",
                }
            ],
        },
    ]


def main():
    out_dir = Path("out_test")
    out_dir.mkdir(exist_ok=True)
    out_file = out_dir / "test_invoices.xlsx"
    invoices = _sample_invoices()
    path = export_invoices_to_excel(invoices, out_file)
    print("Generated:", path)

    # Verificar existencia básica
    p = Path(path)
    if not p.exists():
        print("ERROR: output file not found:", path)
        sys.exit(2)

    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True)
        sheets = wb.sheetnames
        assert "Summary" in sheets, "Summary sheet missing"
        assert "Lines" in sheets, "Lines sheet missing"
        print("OK: sheets present ->", sheets)
        sys.exit(0)
    except Exception as e:
        # Si openpyxl no está presente, aceptamos CSV fallback
        csv_s = out_dir / (out_file.stem + "_Summary.csv")
        csv_l = out_dir / (out_file.stem + "_Lines.csv")
        if csv_s.exists() or csv_l.exists():
            print("OK: openpyxl missing but CSV fallback present")
            sys.exit(0)
        print("ERROR validating workbook:", e)
        sys.exit(3)


if __name__ == "__main__":
    main()
