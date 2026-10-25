"""Exportador específico para facturas: genera un XLSX con hojas Summary y Lines.

Funciona con la estructura de salida consolidada (lista de dicts) que usa claves:
  invoice_id, fecha, proveedor, moneda, subtotal, iva, total, lines[]

Cada línea en `lines` debe tener: origen, descripcion, cantidad, precio_unitario, iva_pct, subtotal, iva_amount, total
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, List

try:
    from openpyxl import Workbook, load_workbook
except Exception:  # pragma: no cover - fallback handled at runtime
    Workbook = None
    load_workbook = None


def _parse_number(s: Any) -> float | None:
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    t = str(s).strip()
    if not t:
        return None
    # quitar símbolos de moneda y espacios
    t = re.sub(r"[€$¥£]|USD|EUR", "", t, flags=re.IGNORECASE)
    t = t.replace(" ", "")

    # manejar formatos 1,234.56 (US) y 1.234,56 (EU)
    if "." in t and "," in t:
        if t.rfind(".") > t.rfind(","):
            # punto = decimal, quitar comas
            t = t.replace(",", "")
        else:
            # coma = decimal, quitar puntos y convertir coma->punto
            t = t.replace(".", "").replace(",", ".")
    elif "," in t:
        # si la parte decimal tiene 2 digitos, asumir coma decimal
        parts = t.split(",")
        if len(parts[-1]) == 2:
            t = t.replace(",", ".")
        else:
            t = t.replace(",", "")

    # dejar solo dígitos, punto y signo
    t = re.sub(r"[^0-9.\-]", "", t)
    try:
        return float(t)
    except Exception:
        return None


def _normalize_date(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    s = str(value).strip()
    if not s:
        return None
    # Intentar parsear con formatos comunes
    fmts = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d"]
    for f in fmts:
        try:
            return datetime.strptime(s, f).date().isoformat()
        except Exception:
            continue
    # fallback: devolver original
    return s


def export_invoices_to_excel(
    invoices: List[dict],
    out_path: str | Path,
    per_invoice_sheets: bool = False,
) -> str:
    """Exporta la lista de facturas a un archivo XLSX.

    - `invoices`: lista de dicts con la estructura normalizada.
    - `out_path`: ruta al fichero a crear (.xlsx). Si es carpeta, se genera con timestamp.
    - `per_invoice_sheets`: si True crea una hoja por factura además de Summary/Lines.
    """
    out = Path(out_path)
    if out.is_dir():
        out = out / f"invoices_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    # Normalizar filas
    summary_rows = []
    lines_rows = []
    for inv in invoices:
        invoice_id = inv.get("invoice_id") or inv.get("id") or ""
        fecha = _normalize_date(inv.get("fecha") or inv.get("date"))
        proveedor = inv.get("proveedor") or inv.get("vendor") or ""
        moneda = inv.get("moneda") or inv.get("currency") or ""
        subtotal = _parse_number(inv.get("subtotal"))
        iva = _parse_number(inv.get("iva") or inv.get("iva_total") or inv.get("tax"))
        total = _parse_number(inv.get("total"))

        summary_rows.append(
            {
                "InvoiceID": invoice_id,
                "Fecha": fecha,
                "Proveedor": proveedor,
                "Moneda": moneda,
                "Subtotal": subtotal,
                "IVA": iva,
                "Total": total,
            }
        )

        for ln in inv.get("lines") or []:
            cantidad = _parse_number(ln.get("cantidad") or ln.get("qty") or ln.get("quantity"))
            precio = _parse_number(ln.get("precio_unitario") or ln.get("unit_price") or ln.get("price"))
            iva_pct = _parse_number(ln.get("iva_pct") or ln.get("iva") or ln.get("tax_pct"))
            subtotal_l = _parse_number(ln.get("subtotal") or ln.get("subtotal_linea"))
            iva_amount = _parse_number(ln.get("iva_amount") or ln.get("tax_amount"))
            total_l = _parse_number(ln.get("total") or ln.get("total_linea"))
            lines_rows.append(
                {
                    "InvoiceID": invoice_id,
                    "Fecha": fecha,
                    "Proveedor": proveedor,
                    "Origen": ln.get("origen") or ln.get("source") or "",
                    "Descripcion": ln.get("descripcion") or ln.get("description") or "",
                    "Cantidad": cantidad,
                    "PrecioUnitario": precio,
                    "Moneda": moneda,
                    "IVA%": iva_pct,
                    "SubtotalLinea": subtotal_l,
                    "IVA_Linea": iva_amount,
                    "TotalLinea": total_l,
                }
            )

    # Intentar usar openpyxl
    if Workbook is None:
        # Fallback: escribir CSVs separados
        import csv

        out_dir = out.parent
        out_dir.mkdir(parents=True, exist_ok=True)
        summary_csv = out_dir / (out.stem + "_Summary.csv")
        lines_csv = out_dir / (out.stem + "_Lines.csv")
        with summary_csv.open("w", newline='', encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=list(summary_rows[0].keys()) if summary_rows else [])
            if summary_rows:
                writer.writeheader()
                writer.writerows(summary_rows)
        with lines_csv.open("w", newline='', encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=list(lines_rows[0].keys()) if lines_rows else [])
            if lines_rows:
                writer.writeheader()
                writer.writerows(lines_rows)
        return str(summary_csv)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    # Summary header
    summary_header = ["InvoiceID", "Fecha", "Proveedor", "Moneda", "Subtotal", "IVA", "Total"]
    ws.append(summary_header)
    for r in summary_rows:
        ws.append([r.get(h) for h in summary_header])

    # Lines sheet
    ws2 = wb.create_sheet("Lines")
    lines_header = [
        "InvoiceID",
        "Fecha",
        "Proveedor",
        "Origen",
        "Descripcion",
        "Cantidad",
        "PrecioUnitario",
        "Moneda",
        "IVA%",
        "SubtotalLinea",
        "IVA_Linea",
        "TotalLinea",
    ]
    ws2.append(lines_header)
    for r in lines_rows:
        ws2.append([r.get(h) for h in lines_header])

    if per_invoice_sheets:
        for inv in invoices:
            sheet_name = str(inv.get("invoice_id") or inv.get("id") or "invoice")
            # sheet names max len 31
            sheet_name = sheet_name[:31]
            wsx = wb.create_sheet(sheet_name)
            # escribir resumen pequeño
            wsx.append(["Key", "Value"])
            wsx.append(["InvoiceID", inv.get("invoice_id") or inv.get("id")])
            wsx.append(["Fecha", _normalize_date(inv.get("fecha") or inv.get("date"))])
            wsx.append(["Proveedor", inv.get("proveedor") or inv.get("vendor")])
            wsx.append(["Subtotal", inv.get("subtotal")])
            wsx.append(["IVA", inv.get("iva")])
            wsx.append(["Total", inv.get("total")])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return str(out)
