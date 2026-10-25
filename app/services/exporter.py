"""Servicio de exportación seguro — TAREA 3.

``DocumentExporter`` convierte el JSON canónico (InvoiceSchema) a
Excel / CSV / Word / TXT con dos garantías:

    * COLUMNAS SELECTIVAS (xlsx/csv): itera sobre la lista ``columnas``
      pedida por el cliente; los campos ``None`` o el string ``"null"``
      se escriben como celda vacía ``""``.
    * SANITIZACIÓN PREVIA (medida de seguridad): antes de exportar se
      detectan datos sensibles — tarjetas de crédito (16 dígitos, con o
      sin separadores) y valores de contraseñas/tokens — y se reemplazan
      por ``[REDACTADO]``. También se bloquean columnas cuyo NOMBRE sea
      sensible (password, token, cvv…).

Convención de columnas planas (xlsx/csv):
    invoice_number, document_type, issue_date, due_date, currency, region,
    supplier, supplier.name, supplier.tax_id, supplier.country,
    customer, customer.name, customer.tax_id, customer.country,
    subtotal, discounts_total, taxes_total, total, taxes, notes,
    full_text_content,
    line.description, line.quantity, line.unit_price, line.tax_rate,
    line.subtotal, line.page

Una fila por ``line_item`` (mínimo una fila); las columnas de nivel
factura se repiten en cada fila (formato ERP-friendly).
"""
from __future__ import annotations

import io
import logging
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

_log = logging.getLogger("ocr.exporter")

# ═══════════════════════════════════════════════════════════════════════════
#  Sanitización (medida de seguridad — TAREA 3.4)
# ═══════════════════════════════════════════════════════════════════════════

REDACTED = "[REDACTADO]"

# Tarjetas de crédito: 16 dígitos con separadores opcionales (4-4-4-4, con
# espacio, guion o ninguno). También un bloque plano de 16 dígitos.
_CARD_PATTERNS = (
    re.compile(r"\b(?:\d{4}[ -]){3}\d{4}\b"),
    re.compile(r"\b\d{16}\b"),
)

# Claves/campos cuyo NOMBRE revela contenido sensible.
_SECRET_NAME_WORDS = (
    "password", "passwd", "pwd", "contrasena", "contraseña", "clave",
    "secret", "token", "api_key", "apikey", "api-key", "cvv", "pin",
)

# Secretos embebidos en texto libre: "password: xyz123", "clave = abc"…
_SECRET_INLINE_RE = re.compile(
    r"(?i)\b(contrase[ñn]a|password|passwd|pwd|clave|secret|token|api[_-]?key|cvv)\b\s*[:=]\s*([^\s;,]+)"
)

# Valores que significan "ausencia de dato" y se exportan como celda vacía.
_NULL_STRINGS = {"", "null", "none"}


def sanitize_value(value: Any) -> str:
    """Limpia UN valor antes de exportarlo.

    * None o "null" → "" (celda vacía, requisito TAREA 3.2).
    * Tarjeta de crédito (16 dígitos) → ``[REDACTADO]``.
    * "password: xxx" / "clave = yyy" → "password: [REDACTADO]".
    """
    if value is None:
        return ""
    text = str(value)
    if text.strip().lower() in _NULL_STRINGS:
        return ""
    for pattern in _CARD_PATTERNS:
        text = pattern.sub(REDACTED, text)
    text = _SECRET_INLINE_RE.sub(lambda m: f"{m.group(1)}: {REDACTED}", text)
    return text


def _is_sensitive_column(column: str) -> bool:
    """True si el NOMBRE de la columna sugiere contenido sensible."""
    lowered = column.lower()
    return any(word in lowered for word in _SECRET_NAME_WORDS)


def _to_decimal(value: Any) -> Decimal | None:
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None


# ═══════════════════════════════════════════════════════════════════════════
#  DocumentExporter
# ═══════════════════════════════════════════════════════════════════════════

DEFAULT_COLUMNS = [
    "invoice_number", "issue_date", "currency",
    "supplier.name", "supplier.tax_id",
    "customer.name",
    "subtotal", "taxes_total", "total",
    "line.description", "line.quantity", "line.unit_price", "line.subtotal",
]

MIME_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain; charset=utf-8",
    "pdf": "application/pdf",
}


class DocumentExporter:
    """Exporta el JSON canónico con columnas selectivas y sanitización."""

    def __init__(self, data: dict, columnas: list[str] | None = None) -> None:
        self.data = data or {}
        # Normaliza la lista de columnas (sin vacíos; vacía → set por defecto).
        requested = [c.strip() for c in (columnas or []) if c and str(c).strip()]
        self.columnas = requested or list(DEFAULT_COLUMNS)
        self._flat = self._flatten_invoice()

    # ── Aplanado del JSON canónico ────────────────────────────────────────
    def _flatten_invoice(self) -> dict[str, Any]:
        """Mapa plano a nivel factura (sin líneas)."""
        d = self.data
        supplier = d.get("supplier") or {}
        customer = d.get("customer") or {}
        validation = d.get("validation") or {}
        taxes = d.get("taxes") or []
        discounts = d.get("discounts") or []

        taxes_total = sum((_to_decimal(t.get("amount")) or Decimal("0") for t in taxes), Decimal("0"))
        discounts_total = sum((_to_decimal(x.get("amount")) or Decimal("0") for x in discounts), Decimal("0"))

        taxes_str = "; ".join(
            f"{t.get('type', '?')}"
            f"{(' ' + str(t['rate']) + '%') if t.get('rate') is not None else ''}"
            f": {t.get('amount', '')}"
            for t in taxes
        )

        return {
            "invoice_number": d.get("invoice_number"),
            "document_type": d.get("document_type"),
            "issue_date": d.get("issue_date"),
            "due_date": d.get("due_date"),
            "currency": d.get("currency"),
            "region": d.get("region"),
            "supplier": supplier.get("name"),
            "supplier.name": supplier.get("name"),
            "supplier.tax_id": supplier.get("tax_id"),
            "supplier.country": supplier.get("country"),
            "customer": customer.get("name"),
            "customer.name": customer.get("name"),
            "customer.tax_id": customer.get("tax_id"),
            "customer.country": customer.get("country"),
            "subtotal": d.get("subtotal"),
            "discounts_total": str(discounts_total) if discounts else "0",
            "taxes_total": str(taxes_total) if taxes else "0",
            "total": d.get("total"),
            "taxes": taxes_str or None,
            "notes": d.get("notes"),
            "full_text_content": d.get("full_text_content"),
            "validation.is_valid": validation.get("is_valid"),
        }

    def _rows(self) -> list[dict[str, str]]:
        """Una fila por línea de detalle; columnas planas + sanitizadas."""
        line_items = self.data.get("line_items") or []
        if not line_items:
            line_items = [{}]  # garantiza al menos una fila

        rows: list[dict[str, str]] = []
        for item in line_items:
            line_flat = {
                "line.description": item.get("description"),
                "line.quantity": item.get("quantity"),
                "line.unit_price": item.get("unit_price"),
                "line.tax_rate": item.get("tax_rate"),
                "line.subtotal": item.get("subtotal"),
                "line.page": item.get("page"),
            }
            row = {}
            for col in self.columnas:
                if _is_sensitive_column(col):
                    # Columna con nombre sensible → nunca exporta el valor real.
                    row[col] = REDACTED
                    _log.warning("Columna sensible solicitada en exportación: '%s'", col)
                else:
                    raw = line_flat.get(col, self._flat.get(col))
                    row[col] = sanitize_value(raw)
            rows.append(row)
        return rows

    # ── Excel (openpyxl) — TAREA 3.1 / 3.2 ────────────────────────────────
    def to_xlsx(self) -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Datos"
        ws.append(self.columnas)
        for row in self._rows():
            ws.append([row[c] for c in self.columnas])
        # Ancho de columna razonable según encabezado y contenido.
        for idx, col in enumerate(self.columnas, start=1):
            width = max([len(col)] + [len(row[col]) for row in self._rows()]) + 2
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(width, 60)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    # ── CSV (pandas) — TAREA 3.1 / 3.2 ────────────────────────────────────
    def to_csv(self) -> bytes:
        import pandas as pd

        df = pd.DataFrame(self._rows(), columns=self.columnas)
        # utf-8-sig: BOM para que Excel en Windows respete acentos/ñ.
        return df.to_csv(index=False).encode("utf-8-sig")

    # ── Word (python-docx) — TAREA 3.3: usa full_text_content ─────────────
    def to_docx(self) -> bytes:
        from docx import Document as DocxDocument

        doc = DocxDocument()
        title = doc.add_heading(
            f"Documento {self._flat.get('invoice_number') or ''}".strip(), level=1
        )
        title.alignment = 1  # centrado

        full_text = self._flat.get("full_text_content")
        if not full_text:
            # Fallback: si la IA no devolvió el texto completo, renderiza el
            # resumen canónico (nunca un documento vacío).
            full_text = self._fallback_text()

        for paragraph in sanitize_value(full_text).splitlines():
            if paragraph.strip():
                doc.add_paragraph(paragraph)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    # ── TXT — TAREA 3.3: usa full_text_content ────────────────────────────
    def to_txt(self) -> bytes:
        full_text = self._flat.get("full_text_content") or self._fallback_text()
        return sanitize_value(full_text).encode("utf-8")

    def _fallback_text(self) -> str:
        """Resumen legible del JSON canónico cuando falta full_text_content."""
        lines = ["(full_text_content no disponible — resumen del JSON canónico)", ""]
        for col in self.columnas:
            if col.startswith("line."):
                continue
            lines.append(f"{col}: {sanitize_value(self._flat.get(col))}")
        lines.append("")
        for i, item in enumerate(self.data.get("line_items") or [], start=1):
            lines.append(
                f"Línea {i}: {sanitize_value(item.get('description'))} x"
                f"{sanitize_value(item.get('quantity'))} = "
                f"{sanitize_value(item.get('subtotal'))}"
            )
        return "\n".join(lines)

    # ── PDF (reportlab) — mantiene el formato del frontend, sanitizado ────
    def to_pdf(self) -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        buffer = io.BytesIO()
        pdf = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
        body = ParagraphStyle("body", fontName="Helvetica", fontSize=9, leading=13)
        story: list = []

        story.append(Paragraph(
            f"<b>Factura {sanitize_value(self._flat.get('invoice_number'))}</b>",
            ParagraphStyle("title", parent=body, fontSize=14, leading=18)))
        story.append(Paragraph(
            f"Fecha: {sanitize_value(self._flat.get('issue_date'))} · "
            f"Moneda: {sanitize_value(self._flat.get('currency'))}", body))
        if self._flat.get("supplier"):
            story.append(Paragraph(
                f"Emisor: {sanitize_value(self._flat.get('supplier.name'))} — "
                f"{sanitize_value(self._flat.get('supplier.tax_id'))}", body))
        if self._flat.get("customer"):
            story.append(Paragraph(
                f"Receptor: {sanitize_value(self._flat.get('customer.name'))} — "
                f"{sanitize_value(self._flat.get('customer.tax_id'))}", body))
        story.append(Spacer(1, 0.5 * cm))

        data = [["Descripción", "Cant.", "P.Unit", "Impuesto", "Subtotal"]]
        for item in self.data.get("line_items") or []:
            data.append([
                sanitize_value(item.get("description")),
                sanitize_value(item.get("quantity")),
                sanitize_value(item.get("unit_price")),
                (f"{sanitize_value(item.get('tax_rate'))}%" if item.get("tax_rate") is not None else ""),
                sanitize_value(item.get("subtotal")),
            ])
        table = Table(data, colWidths=[7 * cm, 2 * cm, 2.5 * cm, 2 * cm, 2.5 * cm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9CA3AF")),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        story.append(table)
        story.append(Spacer(1, 0.4 * cm))
        story.append(Paragraph(
            f"<b>Subtotal:</b> {sanitize_value(self._flat.get('subtotal'))} · "
            f"<b>Impuestos:</b> {sanitize_value(self._flat.get('taxes_total'))} · "
            f"<b>TOTAL: {sanitize_value(self._flat.get('total'))} "
            f"{sanitize_value(self._flat.get('currency'))}</b>", body))

        pdf.build(story)
        return buffer.getvalue()

    # ── Dispatcher ────────────────────────────────────────────────────────
    def export(self, formato: str) -> bytes:
        """Genera los bytes del formato pedido ('xlsx'|'csv'|'docx'|'txt'|'pdf')."""
        fmt = (formato or "").strip().lower()
        exporters = {
            "xlsx": self.to_xlsx,
            "csv": self.to_csv,
            "docx": self.to_docx,
            "txt": self.to_txt,
            "pdf": self.to_pdf,
        }
        exporter = exporters.get(fmt)
        if exporter is None:
            raise ValueError(
                f"Formato de exportación no soportado: '{formato}'. Disponibles: {sorted(exporters)}"
            )
        return exporter()


def export_invoice(inv, formato: str, output_dir: Path, base_name: str) -> Path:
    """Wrapper de compatibilidad: exporta a un archivo en disco.

    (Mantenido para usos existentes; la API REST usa DocumentExporter y
    StreamingResponse directamente.)
    """
    exporter = DocumentExporter(inv.model_dump(mode="json") if hasattr(inv, "model_dump") else inv)
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"{base_name}.{formato.lower()}"
    out_path.write_bytes(exporter.export(formato))
    return out_path
